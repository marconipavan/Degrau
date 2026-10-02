# Fase 6: rotina semanal da turma com planilha fictícia e e-mail simulado (nada sai da máquina).
import os
import pytest
from openpyxl import Workbook
from motor import turma, envio, estado as E

PAINEL = ['Data', 'Aluno', 'Nível', 'Bloco', 'Nota (%)', 'Tempo (min)', 'Limite (min)', 'Status', 'Itens errados', 'Fotos']


def _planilha(caminho, alunos, painel):
    wb = Workbook(); wb.remove(wb.active)
    a = wb.create_sheet('Alunos'); a.append(['Aluno', 'E-mail', 'Nível atual', 'Próximo bloco', 'Entrada (nivelamento)'])
    for nome, email in alunos: a.append([nome, email])
    p = wb.create_sheet('Painel'); p.append(PAINEL)
    for aluno, bloco, nota, status in painel: p.append(['2026-10-05', aluno, 'G1', bloco, nota, 10, 12, status, '', ''])
    wb.create_sheet('Gabarito')
    wb.save(caminho)
    return str(caminho)


class SMTPFalso:
    enviados = []
    def __init__(self, *a, **k): pass
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def login(self, u, s): pass
    def send_message(self, m): SMTPFalso.enviados.append(m)


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    SMTPFalso.enviados = []
    monkeypatch.setattr(envio.smtplib, 'SMTP_SSL', SMTPFalso)
    monkeypatch.setenv('DEGRAU_EMAIL_DE', 'degrau@casd.org'); monkeypatch.setenv('DEGRAU_EMAIL_SENHA', 'x')
    return tmp_path, {'curso': 'geometria-plana-epcar', 'blocos_por_pacote': 3,
                      'link_formulario': 'https://forms.gle/teste', 'pasta_dados': str(tmp_path / 'dados')}


def _por_aluno(resumo): return {r['aluno']: r for r in resumo}


def test_tres_semanas(ambiente):
    tmp, config = ambiente
    alunos = [('Ana Souza', 'ana@x.com'), ('Bruno Lima', 'bruno@x.com'), ('Carla', '')]
    sem = lambda *a: None

    # semana 1: lista nova, Painel vazio -> todos recebem G1 1-3, 4-6, 7-9
    r1 = _por_aluno(turma.preparar(_planilha(tmp / 's1.xlsx', alunos, []), config, hoje='2026-10-03'))
    assert r1['Ana Souza']['pacote']['blocos'] == ['G1 1-3', 'G1 4-6', 'G1 7-9']
    assert not os.path.exists(tmp / 'dados' / 'estados' / 'ana-souza.json')      # nada gravado antes do envio
    turma.enviar_e_gravar(list(r1.values()), config, registro=sem)
    assert [m['To'] for m in SMTPFalso.enviados] == ['ana@x.com', 'bruno@x.com']   # Carla não tem e-mail
    m = SMTPFalso.enviados[0]
    assert [p.get_filename() for p in m.iter_attachments()] == ['G1_pacote01_folhas1-9.pdf']
    assert 'Dia 2: bloco G1 4-6' in m.get_body().get_content() and 'forms.gle/teste' in m.get_body().get_content()
    assert os.path.exists(tmp / 'dados' / 'estados' / 'ana-souza.json')
    assert not os.path.exists(tmp / 'dados' / 'estados' / 'carla.json')            # sem envio, sem estado

    # semana 2: Ana reenviou o G1 4-6 (vale o mais recente); Bruno não mandou o G1 7-9
    painel = [('Ana Souza', 'G1 1-3', 100, 'Domínio'), ('Ana Souza', 'G1 4-6', 70, 'Repetir'),
              ('Ana Souza', 'G1 4-6', 95, 'Domínio'), ('Ana Souza', 'G1 7-9', 92, 'Domínio'),
              ('Bruno Lima', 'G1 1-3', 90, 'Domínio'), ('Bruno Lima', 'G1 4-6', 91, 'Domínio'),
              ('Bruno Lima', 'G1 99', '', 'ERRO: sem gabarito')]
    SMTPFalso.enviados = []
    r2 = _por_aluno(turma.preparar(_planilha(tmp / 's2.xlsx', alunos, painel), config, hoje='2026-10-07'))
    assert r2['Ana Souza']['pacote']['blocos'] == ['G1 10-12', 'G1 13-15', 'G1 16-18']
    assert r2['Bruno Lima']['pendentes'] == ['G1 7-9'] and r2['Bruno Lima']['pacote'] is None
    assert r2['Carla']['pacote']['blocos'] == ['G1 1-3', 'G1 4-6', 'G1 7-9']
    assert 'PENDENTE   G1 7-9' in turma.texto_resumo(list(r2.values()))
    turma.enviar_e_gravar(list(r2.values()), config, registro=sem)
    assunto = {m['To']: m['Subject'] for m in SMTPFalso.enviados}
    assert assunto['bruno@x.com'] == 'Degrau — falta enviar: G1 7-9'
    assert not list(SMTPFalso.enviados[1].iter_attachments())                     # lembrete sem anexo

    # semana 3: Bruno repete o G1 7-9, mas a versão 2 dessas folhas não existe -> aviso, nada enviado
    painel.append(('Bruno Lima', 'G1 7-9', 60, 'Repetir'))
    SMTPFalso.enviados = []
    r3 = _por_aluno(turma.preparar(_planilha(tmp / 's3.xlsx', alunos, painel), config, hoje='2026-10-10'))
    assert 'faltam folhas' in r3['Bruno Lima']['erro'] and 'G1/7.v2.yaml' in r3['Bruno Lima']['erro']
    turma.enviar_e_gravar(list(r3.values()), config, registro=sem)
    assert 'bruno@x.com' not in [m['To'] for m in SMTPFalso.enviados]
    bruno = E.carregar(tmp / 'dados' / 'estados' / 'bruno-lima.json')
    assert [u['status'] for u in bruno['unidades']] == ['dominio', 'dominio', 'pendente']   # repetição não gravada


def test_planilha_sem_aba(ambiente, tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); wb.save(tmp_path / 'x.xlsx')
    with pytest.raises(turma.ErroEspecificacao, match='não tem a aba Alunos'):
        turma.ler_planilha(str(tmp_path / 'x.xlsx'))
