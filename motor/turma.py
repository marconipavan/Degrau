# Degrau — rotina semanal da turma (CASD): planilha baixada -> resultados -> próximos pacotes -> e-mails.
#
# Quem opera baixa a planilha do Google em Excel e roda: degrau semana planilha.xlsx
#  1. lê a aba Alunos (nome, e-mail) e a aba Painel (resultados já corrigidos pelo Apps Script);
#  2. registra no estado de cada aluno o resultado de cada bloco pendente (vale o envio mais recente);
#  3. gera o próximo pacote de quem não tem bloco pendente; quem tem recebe um lembrete;
#  4. mostra o resumo e, com a confirmação, envia os e-mails.
# Estados e pacotes ficam em <pasta_dados>/ (fora do git). Todo aluno novo começa na primeira folha do curso.
import os, re, unicodedata
from datetime import date
import yaml
from openpyxl import load_workbook
from .especificacao import ErroEspecificacao, RAIZ, carregar_curso
from . import estado as E, gerador

STATUS = {'Domínio': 'dominio', 'Repetir': 'repetir'}


def carregar_config(caminho=None):
    caminho = caminho or os.path.join(RAIZ, 'config.local.yaml')
    if not os.path.exists(caminho):
        raise ErroEspecificacao(f'{caminho} não existe: copie config.exemplo.yaml para config.local.yaml e ajuste')
    with open(caminho, encoding='utf-8') as f:
        c = yaml.safe_load(f)
    for k in ('curso', 'blocos_por_pacote', 'link_formulario', 'pasta_dados'):
        if k not in c: raise ErroEspecificacao(f'{caminho}: falta {k}')
    if not os.path.isabs(c['pasta_dados']): c['pasta_dados'] = os.path.join(RAIZ, c['pasta_dados'])
    return c


def _aba(wb, nome, caminho):
    if nome not in wb.sheetnames:
        raise ErroEspecificacao(f'{caminho}: a planilha não tem a aba {nome} (baixe a planilha inteira em Excel)')
    linhas = [list(l) for l in wb[nome].iter_rows(values_only=True)]
    if not linhas: return []
    cab = [str(c or '').strip() for c in linhas[0]]
    return [dict(zip(cab, l)) for l in linhas[1:] if any(v not in (None, '') for v in l)]


def ler_planilha(caminho):
    """(alunos [{'nome', 'email'}], painel [linhas do Painel na ordem em que chegaram])"""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    alunos = []
    for l in _aba(wb, 'Alunos', caminho):
        nome = str(l.get('Aluno') or '').strip()
        if nome: alunos.append({'nome': nome, 'email': str(l.get('E-mail') or '').strip()})
    return alunos, _aba(wb, 'Painel', caminho)


def _arquivo(nome):
    """nome do aluno -> nome de arquivo sem acento nem espaço"""
    s = unicodedata.normalize('NFD', nome).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-') or 'aluno'


def _num(v):
    try: return float(v)
    except (TypeError, ValueError): return None


def preparar(planilha, config, hoje=None):
    """Lê a planilha e gera os pacotes. Devolve o resumo por aluno. Nada é enviado nem gravado no estado aqui:
    o estado de cada aluno só é gravado depois que o e-mail dele sai (enviar_e_gravar)."""
    curso = carregar_curso(config['curso'])
    inicio = f'{curso["niveis"][0]["codigo"]} 1'
    alunos, painel = ler_planilha(planilha)
    pasta = config['pasta_dados']
    resumo = []
    for a in alunos:
        r = {'aluno': a['nome'], 'email': a['email'], 'registrados': [], 'pendentes': [], 'pacote': None, 'erro': None}
        caminho = os.path.join(pasta, 'estados', _arquivo(a['nome']) + '.json')
        if os.path.exists(caminho):
            est = E.carregar(caminho)
        else:
            os.makedirs(os.path.dirname(caminho), exist_ok=True)
            est = {'curso': config['curso'], 'aluno': a['nome'], 'inicio': inicio, 'unidades': [],
                   'vocabulario': {}, 'estruturas': {}, 'erros_recorrentes': []}
        # resultados do Painel: vale o envio mais recente de cada bloco
        ultimo = {}
        for l in painel:
            if str(l.get('Aluno') or '').strip() == a['nome'] and str(l.get('Status')) in STATUS:
                ultimo[str(l.get('Bloco')).strip()] = l
        for u in list(gerador.pendentes(est)):
            l = ultimo.get(u['codigo'])
            if l is None: continue
            gerador.registrar_resultado(est, u['codigo'], _num(l.get('Nota (%)')), _num(l.get('Tempo (min)')),
                                        STATUS[str(l['Status'])])
            r['registrados'].append(f'{u["codigo"]}: {l["Status"]} ({l.get("Nota (%)")}%)')
        r['pendentes'] = [u['codigo'] for u in gerador.pendentes(est)]
        if not r['pendentes']:
            try:
                arquivos, _ = gerador.proximo(est, dias=config['blocos_por_pacote'], hoje=hoje,
                                              saida=os.path.join(pasta, 'pacotes', _arquivo(a['nome'])))
                n = est['unidades'][-1]['pacote']
                r['pacote'] = {'pdf': arquivos[0], 'blocos': [u['codigo'] for u in est['unidades'] if u['pacote'] == n]}
            except ErroEspecificacao as e:
                r['erro'] = str(e).replace(RAIZ + os.sep, '')
        r['_estado'], r['_caminho'] = est, caminho
        resumo.append(r)
    return resumo


def enviar_e_gravar(resumo, config, enviar=True, registro=print):
    """Envia o pacote (ou o lembrete) de cada aluno e só então grava o estado dele. Com enviar=False, só grava."""
    from . import envio
    for r in resumo:
        if r['erro'] and not r['pendentes']:
            registro(f'  {r["aluno"]}: não enviado ({r["erro"].splitlines()[0]})'); continue
        if enviar:
            if not r['email']:
                registro(f'  {r["aluno"]}: sem e-mail na aba Alunos, nada enviado'); continue
            try:
                envio.enviar(envio.montar_turma(r, config))
            except ErroEspecificacao as e:
                registro(f'  {r["aluno"]}: falhou ({e}); estado não gravado'); continue
        E.salvar(r['_caminho'], r['_estado'])
        registro(f'  {r["aluno"]}: {"pacote" if r["pacote"] else "lembrete"} {"enviado" if enviar else "gravado"}')


def texto_resumo(resumo):
    linhas = []
    for r in resumo:
        linhas.append(f'{r["aluno"]} <{r["email"] or "sem e-mail"}>')
        for x in r['registrados']: linhas.append(f'    resultado  {x}')
        if r['pendentes']: linhas.append(f'    PENDENTE   {", ".join(r["pendentes"])} (recebe lembrete, não recebe pacote novo)')
        if r['pacote']: linhas.append(f'    pacote     {", ".join(r["pacote"]["blocos"])}  ->  {os.path.basename(r["pacote"]["pdf"])}')
        if r['erro']: linhas.append(f'    ERRO       {r["erro"]}')
    return '\n'.join(linhas)
